
from pathlib import Path
from datetime import datetime
import re
import json
import pymupdf

from bs4 import BeautifulSoup
from docx import Document


# ============================================================
# CHEMINS ET FORMATS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent.parent

CORPUS_DIR = BASE_DIR / "corpus_complet"
DATA_CLEAN_DIR = BASE_DIR / "data_clean"

FORMATS_ACCEPTES = {".pdf", ".html", ".htm", ".docx"}

DATA_CLEAN_DIR.mkdir(parents=True, exist_ok=True)


# ============================================================
# RÉCUPÉRATION DES DOCUMENTS
# ============================================================

documents = sorted([
    fichier
    for fichier in CORPUS_DIR.iterdir()
    if fichier.is_file()
    and fichier.suffix.lower() in FORMATS_ACCEPTES
])

print(f"Nombre de documents trouvés : {len(documents)}")


# ============================================================
# EXTRACTION PDF AVEC NUMÉROS DE PAGE
# ============================================================

def extraire_pdf(chemin, avec_pages=False):

    pages = []

    with pymupdf.open(chemin) as document:

        for numero, page in enumerate(document, start=1):

            blocs = page.get_text("blocks", sort=True)

            textes = [
                bloc[4].strip()
                for bloc in blocs
                if len(bloc) > 4
                and isinstance(bloc[4], str)
                and bloc[4].strip()
            ]

            texte_page = "\n\n".join(textes)

            pages.append({
                "numero_page": numero,
                "texte": texte_page
            })

    if avec_pages:
        return pages

    return "\n\n".join(
        page["texte"]
        for page in pages
    )


# ============================================================
# EXTRACTION HTML
# ============================================================

def extraire_html(chemin):

    with open(chemin, "r", encoding="utf-8", errors="replace") as fichier:
        contenu = fichier.read()

    soup = BeautifulSoup(contenu, "html.parser")

    # Supprimer les éléments non documentaires
    for element in soup([
        "script",
        "style",
        "noscript",
        "nav",
        "footer"
    ]):
        element.decompose()

    return soup.get_text(separator="\n", strip=True)


# ============================================================
# EXTRACTION DOCX
# ============================================================

def extraire_docx(chemin):

    document = Document(chemin)
    contenus = []

    # Extraction des paragraphes
    for paragraphe in document.paragraphs:

        texte = paragraphe.text.strip()

        if texte:
            contenus.append(texte)

    # Extraction des tableaux
    for tableau in document.tables:

        for ligne in tableau.rows:

            cellules = []

            for cellule in ligne.cells:

                texte = cellule.text.strip()

                if texte:
                    cellules.append(texte)

            if cellules:
                contenus.append(" | ".join(cellules))

    return "\n".join(contenus)


# ============================================================
# NETTOYAGE DU TEXTE
# ============================================================

def nettoyer_texte(texte):

    # Uniformiser les retours à la ligne
    texte = texte.replace("\r\n", "\n")
    texte = texte.replace("\r", "\n")

    # Corriger les mots coupés en fin de ligne
    texte = re.sub(
        r"(\w)-\n(\w)",
        r"\1\2",
        texte
    )

    # Identifier temporairement les paragraphes
    texte = re.sub(
        r"\n\s*\n+",
        " <PARAGRAPHE> ",
        texte
    )

    # Fusionner les retours à la ligne simples
    texte = re.sub(
        r"\s*\n\s*",
        " ",
        texte
    )

    # Restaurer les paragraphes
    texte = texte.replace(
        "<PARAGRAPHE>",
        "\n\n"
    )

    # Uniformiser les espaces
    texte = re.sub(
        r"[ \t]+",
        " ",
        texte
    )

    # Supprimer les espaces avant la ponctuation
    texte = re.sub(
        r"\s+([,.;:!?])",
        r"\1",
        texte
    )

    # Ajouter les espaces manquants après la ponctuation
    texte = re.sub(
        r"([,;:!?])(?=[A-Za-zÀ-ÿ])",
        r"\1 ",
        texte
    )

    # Ajouter un espace après un point suivi d'une majuscule
    texte = re.sub(
        r"\.(?=[A-ZÀ-Ý])",
        ". ",
        texte
    )

    # Nettoyer les espaces autour des parenthèses
    texte = re.sub(
        r"\(\s+",
        "(",
        texte
    )

    texte = re.sub(
        r"\s+\)",
        ")",
        texte
    )

    # Limiter les retours à la ligne
    texte = re.sub(
        r"\n{3,}",
        "\n\n",
        texte
    )

    return texte.strip()


# ============================================================
# EXTRACTION AUTOMATIQUE SELON LE FORMAT
# ============================================================

def extraire_document(chemin):

    extension = chemin.suffix.lower()

    if extension == ".pdf":
        return extraire_pdf(chemin)

    elif extension in {".html", ".htm"}:
        return extraire_html(chemin)

    elif extension == ".docx":
        return extraire_docx(chemin)

    return ""


# ============================================================
# EXPORT DE TOUS LES DOCUMENTS EN JSON
# ============================================================

print("\n--- DÉBUT DU TRAITEMENT ---")

nombre_reussis = 0
nombre_erreurs = 0

for fichier in documents:

    try:

        # ----------------------------------------------------
        # 1. EXTRACTION DU TEXTE
        # ----------------------------------------------------

        texte_brut = extraire_document(fichier)

        # ----------------------------------------------------
        # 2. RÉCUPÉRATION DES PAGES PDF
        # ----------------------------------------------------

        pages_nettoyees = []

        if fichier.suffix.lower() == ".pdf":

            pages_brutes = extraire_pdf(
                fichier,
                avec_pages=True
            )

            pages_nettoyees = [
                {
                    "numero_page": page["numero_page"],
                    "texte": nettoyer_texte(page["texte"])
                }
                for page in pages_brutes
            ]

        # ----------------------------------------------------
        # 3. NETTOYAGE DU TEXTE COMPLET
        # ----------------------------------------------------

        texte_nettoye = nettoyer_texte(texte_brut)

        # ----------------------------------------------------
        # 4. VÉRIFICATION
        # ----------------------------------------------------

        if not texte_nettoye:
            raise ValueError(
                "Aucun texte extrait du document"
            )

        # ----------------------------------------------------
        # 5. DATE DU FICHIER
        # ----------------------------------------------------

        date_fichier = datetime.fromtimestamp(
            fichier.stat().st_mtime
        ).strftime("%Y-%m-%d %H:%M:%S")

        # ----------------------------------------------------
        # 6. MÉTADONNÉES + TEXTE + PAGES
        # ----------------------------------------------------

        donnees = {
            "nom_fichier": fichier.name,
            "format": fichier.suffix.lower(),
            "date": date_fichier,
            "nombre_caracteres_brut": len(texte_brut),
            "nombre_caracteres_nettoye": len(texte_nettoye),
            "texte": texte_nettoye,
            "pages": pages_nettoyees
        }

        # ----------------------------------------------------
        # 7. CRÉATION DU NOM DU JSON
        # ----------------------------------------------------

        nom_sans_extension = fichier.name

        # Retirer les extensions répétées
        while (
            Path(nom_sans_extension).suffix.lower()
            in FORMATS_ACCEPTES
        ):
            nom_sans_extension = Path(
                nom_sans_extension
            ).stem

        chemin_json = (
            DATA_CLEAN_DIR /
            f"{nom_sans_extension}.json"
        )

        # ----------------------------------------------------
        # 8. ENREGISTREMENT DU JSON
        # ----------------------------------------------------

        with open(
            chemin_json,
            "w",
            encoding="utf-8"
        ) as fichier_json:

            json.dump(
                donnees,
                fichier_json,
                ensure_ascii=False,
                indent=4
            )

        # ----------------------------------------------------
        # 9. DOCUMENT TRAITÉ AVEC SUCCÈS
        # ----------------------------------------------------

        nombre_reussis += 1

        print(
            f"[OK] {fichier.name} "
            f"-> {chemin_json.name}"
        )

    except Exception as erreur:

        nombre_erreurs += 1

        print(
            f"[ERREUR] {fichier.name} "
            f"-> {erreur}"
        )


# ============================================================
# BILAN FINAL
# ============================================================

print("\n--- BILAN EXPORT ---")

print(f"Documents trouvés : {len(documents)}")
print(f"JSON créés : {nombre_reussis}")
print(f"Erreurs : {nombre_erreurs}")


# ============================================================
# VÉRIFICATION DES PAGES D'UN PDF
# ============================================================

fichier_test = (
    CORPUS_DIR /
    "18_digital_twins_clinical_oncology.pdf"
)

if fichier_test.exists():

    pages_test = extraire_pdf(
        fichier_test,
        avec_pages=True
    )

    print("\n--- TEST EXTRACTION DES PAGES ---")

    print(
        f"Nombre de pages : {len(pages_test)}"
    )

    for page in pages_test[:3]:

        print(
            f"\nPage {page['numero_page']} :"
        )

        print(
            page["texte"][:200]
        )

